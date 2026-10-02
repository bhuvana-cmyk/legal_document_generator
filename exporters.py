
from io import BytesIO
from docx import Document
from docx.shared import Pt
from fpdf import FPDF

def make_txt(text):
    return text.encode("utf-8")

def make_docx(text):
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    for line in text.splitlines():
        if line.strip():
            doc.add_paragraph(line)
        else:
            doc.add_paragraph("")

    output = BytesIO()
    doc.save(output)
    return output.getvalue()

def make_pdf(text):
    pdf = FPDF()
    pdf.set_auto_page_break(True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", size=11)

    for line in text.splitlines():
        safe_line = line.encode(
            "latin-1", "replace"
        ).decode("latin-1")

        if safe_line.strip():
            pdf.multi_cell(0, 7, safe_line)
        else:
            pdf.ln(4)

    return bytes(pdf.output())
  
